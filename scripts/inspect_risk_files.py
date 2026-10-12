import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

risk_dir = r"F:\Mr Hesham\TMG\RISK"
files = [
    "TMG_Project_Companies_Contacts_One_Sheet.xlsx",
    "TMG_Project_Teams_LinkedIn_People_AR.xlsx",
    "tmg_swot_analysis.xlsx"
]

for filename in files:
    filepath = os.path.join(risk_dir, filename)
    print("=" * 60)
    print(f"FILE: {filename}")
    wb = openpyxl.load_workbook(filepath, data_only=True)
    print(f"Sheet names: {wb.sheetnames}")
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        print(f"\n--- Sheet: {sheetname} (max_row={sheet.max_row}, max_column={sheet.max_column}) ---")
        headers = [sheet.cell(row=1, column=c).value for c in range(1, min(sheet.max_column + 1, 15))]
        print(f"Headers: {headers}")
        print("Sample Row 2:")
        row2 = [sheet.cell(row=2, column=c).value for c in range(1, min(sheet.max_column + 1, 15))]
        print(row2)
