import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')
risk_dir = r"F:\Mr Hesham\TMG\RISK"

def inspect_sheet_rows(filename, sheetname, n_rows=8):
    filepath = os.path.join(risk_dir, filename)
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb[sheetname]
    print(f"\n==================== {filename} -> {sheetname} ====================")
    for r in range(1, min(sheet.max_row + 1, n_rows + 1)):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, min(sheet.max_column + 1, 15))]
        # print non-empty row
        if any(v is not None for v in row_vals):
            print(f"Row {r:2d}: {row_vals}")

inspect_sheet_rows("TMG_Project_Companies_Contacts_One_Sheet.xlsx", "دليل الاتصال")
inspect_sheet_rows("TMG_Project_Teams_LinkedIn_People_AR.xlsx", "فرق المشاريع وLinkedIn")
inspect_sheet_rows("tmg_swot_analysis.xlsx", "SWOT Overview & Scoring")
inspect_sheet_rows("tmg_swot_analysis.xlsx", "Strategic TOWS Matrix")
