import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')
risk_dir = r"F:\Mr Hesham\TMG\RISK"

filename = "Balance sheet and   .xlsx"
filepath = os.path.join(risk_dir, filename)

if os.path.exists(filepath):
    print(f"Loading {filepath}...")
    wb = openpyxl.load_workbook(filepath, data_only=True)
    print("Sheet names:", wb.sheetnames)
    for name in wb.sheetnames:
        sh = wb[name]
        print(f"\n--- Sheet: {name} (max_row={sh.max_row}, max_column={sh.max_column}) ---")
        for r in range(1, min(sh.max_row + 1, 15)):
            row_vals = [sh.cell(row=r, column=c).value for c in range(1, min(sh.max_column + 1, 12))]
            if any(v is not None for v in row_vals):
                print(f"R{r}: {row_vals}")
else:
    print(f"File not found: {filepath}")
