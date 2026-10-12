import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'F:\Mr Hesham\TMG\TMG_Holding_Corporate_Risk_Analysis_AR.docx')

print("Document Sections:", len(doc.sections))
print("Paragraphs count:", len(doc.paragraphs))
print("Tables count:", len(doc.tables))

for i, p in enumerate(doc.paragraphs[:40]):
    if p.text.strip():
        print(f"P{i+1} [{p.style.name}]: {p.text[:120]}")

print("\n--- TABLES SUMMARY ---")
for t_idx, table in enumerate(doc.tables):
    rows = len(table.rows)
    cols = len(table.columns) if rows > 0 else 0
    header = [c.text.strip().replace('\n', ' ') for c in table.rows[0].cells] if rows > 0 else []
    print(f"Table {t_idx+1}: {rows} rows x {cols} cols | Header: {header[:4]}")
