import openpyxl
import json
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')
risk_dir = r"F:\Mr Hesham\TMG\RISK"

# 1. Parse Companies
wb_comp = openpyxl.load_workbook(os.path.join(risk_dir, "TMG_Project_Companies_Contacts_One_Sheet.xlsx"), data_only=True)
sheet_comp = wb_comp["دليل الاتصال"]
companies = []
for r in range(5, sheet_comp.max_row + 1):
    c_id = sheet_comp.cell(row=r, column=1).value
    if c_id is None:
        continue
    name = sheet_comp.cell(row=r, column=2).value or ""
    role = sheet_comp.cell(row=r, column=3).value or ""
    project = sheet_comp.cell(row=r, column=4).value or ""
    relation = sheet_comp.cell(row=r, column=5).value or ""
    phone = sheet_comp.cell(row=r, column=6).value or ""
    phone_notes = sheet_comp.cell(row=r, column=7).value or ""
    email = sheet_comp.cell(row=r, column=8).value or ""
    website = sheet_comp.cell(row=r, column=9).value or ""
    rel_source = sheet_comp.cell(row=r, column=10).value or ""
    phone_source = sheet_comp.cell(row=r, column=11).value or ""
    evidence_level = sheet_comp.cell(row=r, column=12).value or ""
    confidence = sheet_comp.cell(row=r, column=13).value or ""
    notes = sheet_comp.cell(row=r, column=14).value or ""
    
    companies.append({
        "id": int(c_id) if isinstance(c_id, (int, float)) else str(c_id),
        "name": str(name).strip(),
        "role": str(role).strip(),
        "project": str(project).strip(),
        "relation": str(relation).strip(),
        "phone": str(phone).strip(),
        "phone_notes": str(phone_notes).strip(),
        "email": str(email).strip(),
        "website": str(website).strip(),
        "rel_source": str(rel_source).strip(),
        "phone_source": str(phone_source).strip(),
        "evidence_level": str(evidence_level).strip(),
        "confidence": str(confidence).strip(),
        "notes": str(notes).strip()
    })

print(f"Extracted {len(companies)} companies.")

# 2. Parse LinkedIn Teams
wb_team = openpyxl.load_workbook(os.path.join(risk_dir, "TMG_Project_Teams_LinkedIn_People_AR.xlsx"), data_only=True)
sheet_team = wb_team["فرق المشاريع وLinkedIn"]
teams = []
for r in range(4, sheet_team.max_row + 1):
    t_id = sheet_team.cell(row=r, column=1).value
    if t_id is None:
        continue
    project = sheet_team.cell(row=r, column=2).value or ""
    company = sheet_team.cell(row=r, column=3).value or ""
    name = sheet_team.cell(row=r, column=4).value or ""
    role = sheet_team.cell(row=r, column=5).value or ""
    linkedin = sheet_team.cell(row=r, column=6).value or ""
    phone = sheet_team.cell(row=r, column=7).value or ""
    evidence_level = sheet_team.cell(row=r, column=8).value or ""
    source = sheet_team.cell(row=r, column=9).value or ""
    notes = sheet_team.cell(row=r, column=10).value or ""
    
    teams.append({
        "id": int(t_id) if isinstance(t_id, (int, float)) else str(t_id),
        "project": str(project).strip(),
        "company": str(company).strip(),
        "name": str(name).strip(),
        "role": str(role).strip(),
        "linkedin": str(linkedin).strip(),
        "phone": str(phone).strip(),
        "evidence_level": str(evidence_level).strip(),
        "source": str(source).strip(),
        "notes": str(notes).strip()
    })

print(f"Extracted {len(teams)} key personnel profiles.")

# 3. Parse SWOT Details
wb_swot = openpyxl.load_workbook(os.path.join(risk_dir, "tmg_swot_analysis.xlsx"), data_only=True)

def parse_swot_sheet(sheetname):
    sh = wb_swot[sheetname]
    items = []
    for r in range(5, sh.max_row + 1):
        s_id = sh.cell(row=r, column=1).value
        if not s_id:
            continue
        factor = sh.cell(row=r, column=2).value or ""
        details = sh.cell(row=r, column=3).value or ""
        quant_data = sh.cell(row=r, column=4).value or ""
        impact = sh.cell(row=r, column=5).value or ""
        source = sh.cell(row=r, column=6).value or ""
        items.append({
            "id": str(s_id).strip(),
            "factor": str(factor).strip(),
            "details": str(details).strip(),
            "data": str(quant_data).strip(),
            "impact": str(impact).strip(),
            "source": str(source).strip()
        })
    return items

strengths = parse_swot_sheet("Detailed Strengths")
weaknesses = parse_swot_sheet("Detailed Weaknesses")
opportunities = parse_swot_sheet("Detailed Opportunities")
threats = parse_swot_sheet("Detailed Threats")

print(f"Parsed SWOT: {len(strengths)} Strengths, {len(weaknesses)} Weaknesses, {len(opportunities)} Opportunities, {len(threats)} Threats.")

# 4. Parse TOWS Strategies
sh_tows = wb_swot["Strategic TOWS Matrix"]
tows = {"SO": [], "ST": [], "WO": [], "WT": []}
current_cat = None
for r in range(4, sh_tows.max_row + 1):
    c1 = sh_tows.cell(row=r, column=1).value
    if c1 and "SO STRATEGIES" in str(c1):
        current_cat = "SO"
        continue
    elif c1 and "ST STRATEGIES" in str(c1):
        current_cat = "ST"
        continue
    elif c1 and "WO STRATEGIES" in str(c1):
        current_cat = "WO"
        continue
    elif c1 and "WT STRATEGIES" in str(c1):
        current_cat = "WT"
        continue
    elif c1 == "Strategy ID" or not c1:
        continue
    
    if current_cat and str(c1).startswith(current_cat):
        strat_id = str(c1).strip()
        name = sh_tows.cell(row=r, column=2).value or ""
        roadmap = sh_tows.cell(row=r, column=3).value or ""
        priority = sh_tows.cell(row=r, column=7).value or "High Priority"
        tows[current_cat].append({
            "id": strat_id,
            "name": str(name).strip(),
            "roadmap": str(roadmap).strip(),
            "priority": str(priority).strip()
        })

for k, v in tows.items():
    print(f"TOWS {k}: {len(v)} strategies.")
