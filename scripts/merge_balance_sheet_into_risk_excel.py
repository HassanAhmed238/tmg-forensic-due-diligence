# -*- coding: utf-8 -*-
"""
Merge Audited Balance Sheet, Income Statement, Common Size, and Financial Ratios
directly into TMG_Key_Risk_Assessments_and_Weibull_Analysis.xlsx
so that the workbook itself contains the complete TMG Balance Sheet with formulas.
"""
import os, sys
import openpyxl
from copy import copy

sys.stdout.reconfigure(encoding='utf-8')

risk_file = r"F:\Mr Hesham\TMG\TMG_Key_Risk_Assessments_and_Weibull_Analysis.xlsx"
fin_file = r"F:\Mr Hesham\TMG\RISK\Balance sheet and TMG Real Data.xlsx"

print(f"Loading risk workbook: {risk_file}")
wb_risk = openpyxl.load_workbook(risk_file)

print(f"Loading financial workbook: {fin_file}")
wb_fin = openpyxl.load_workbook(fin_file, data_only=False)

sheets_to_copy = [
    'Financial statement',
    'Source & uses analysis ',
    'Common size analysis',
    'Ratios & CFO Scorecard'
]

for sheet_name in sheets_to_copy:
    if sheet_name in wb_risk.sheetnames:
        print(f"Sheet {sheet_name} already exists in risk wb, removing old version...")
        wb_risk.remove(wb_risk[sheet_name])
    
    src_ws = wb_fin[sheet_name]
    target_title = sheet_name.strip()
    if len(target_title) > 31:
        target_title = target_title[:31]
    
    dest_ws = wb_risk.create_sheet(title=target_title)
    dest_ws.views.sheetView[0].showGridLines = True
    
    # Copy column dimensions
    for col_letter, col_dim in src_ws.column_dimensions.items():
        dest_ws.column_dimensions[col_letter].width = col_dim.width
        
    # Copy row dimensions
    for row_idx, row_dim in src_ws.row_dimensions.items():
        dest_ws.row_dimensions[row_idx].height = row_dim.height
        
    # Copy cells first (before merging)
    for r in range(1, src_ws.max_row + 1):
        for c in range(1, src_ws.max_column + 1):
            src_cell = src_ws.cell(r, c)
            dest_cell = dest_ws.cell(r, c)
            dest_cell.value = src_cell.value
            
            if src_cell.has_style:
                dest_cell.font = copy(src_cell.font)
                dest_cell.border = copy(src_cell.border)
                dest_cell.fill = copy(src_cell.fill)
                dest_cell.number_format = copy(src_cell.number_format)
                dest_cell.protection = copy(src_cell.protection)
                dest_cell.alignment = copy(src_cell.alignment)

    # Now apply merged cells
    for merge_range in src_ws.merged_cells.ranges:
        dest_ws.merge_cells(str(merge_range))
                
    print(f"Successfully copied sheet: {target_title}")

# Save updated workbook
wb_risk.save(risk_file)
print(f"Successfully saved updated: {risk_file}")

# Also save master copy
master_copy = r"F:\Mr Hesham\TMG\TMG_Master_Forensic_Audit_Model.xlsx"
wb_risk.save(master_copy)
print(f"Successfully saved master copy: {master_copy}")
