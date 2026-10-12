import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

for path in [
    r'F:\Mr Hesham\TMG\TMG_Holding_Corporate_Risk_Analysis.docx',
    r'F:\Mr Hesham\TMG\TMG_Holding_Corporate_Risk_Analysis_AR.docx'
]:
    doc = docx.Document(path)
    print(f"\n=======================================================")
    print(f"FILE: {path}")
    print(f"Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
    
    headings = [p.text.strip() for p in doc.paragraphs if p.text.strip() and p.style.name.startswith('Heading')]
    print(f"Headings count: {len(headings)}")
    for h in headings[:15]:
        print(f"  - {h}")
    
    # Check for specific search terms
    full_text = " ".join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            full_text += " " + " ".join([cell.text.strip() for cell in row.cells])
            
    terms = [
        'شحاتة', 'Shehata', 'حسبو', 'Hasabo', 'المصري اليوم', 'Al-Masry',
        'مدى مصر', 'Mada Masr', 'المنصة', 'Manassa', 'أفراد أمن', 'Security',
        'SouthMED', 'Banan', 'Noor', 'Atrium', 'Concrete Plus', '15.79', '44', 'Rawasy',
        'Luxor', 'الأقصر', 'EGOTH', 'إيجوث', 'CCC', 'Shaker', 'Hill', 'Knauf'
    ]
    print("\nTerm Presence:")
    for t in terms:
        print(f"  - '{t}': {t.lower() in full_text.lower()}")
